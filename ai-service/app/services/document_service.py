import pdfplumber
from docx import Document
import chardet
import uuid
from datetime import datetime
from pathlib import Path
from typing import Dict
from pypdf import PdfReader

from app.utils.text_processing import clean_text, chunk_text
from app.config import settings

class DocumentService:
    """Handle PDF, DOCX, and TXT file extraction"""
    
    def __init__(self, chunk_size: int = settings.chunk_size, chunk_overlap: int = settings.chunk_overlap):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
    
    async def extract_pdf(self, file_path: str) -> Dict:
        """Extract text from PDF"""
        try:
            text = ""
            total_pages = 0
            
            # Try pdfplumber first (better for complex layouts)
            try:
                with pdfplumber.open(file_path) as pdf:
                    total_pages = len(pdf.pages)
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += f"{page_text}\n"
            except Exception as e:
                # Fallback to pypdf
                print(f"pdfplumber failed, falling back to pypdf: {e}")
                reader = PdfReader(file_path)
                total_pages = len(reader.pages)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += f"{page_text}\n"
                        
            if not text.strip():
                return {"success": False, "error": "Could not extract text from PDF. The document might be scanned or image-based."}
                
            return {
                "success": True,
                "text": text,
                "page_count": total_pages
            }
        except Exception as e:
            return {"success": False, "error": str(e)}
    
    async def extract_docx(self, file_path: str) -> Dict:
        """Extract text from DOCX"""
        try:
            doc = Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            
            if not text.strip():
                return {"success": False, "error": "Could not extract text from DOCX. Document might be empty."}
                
            return {
                "success": True,
                "text": text,
                "paragraphs": len(doc.paragraphs)
            }
        except Exception as e:
            return {"success": False, "error": str(e)}
    
    async def extract_txt(self, file_path: str) -> Dict:
        """Extract text from TXT"""
        try:
            with open(file_path, 'rb') as f:
                raw_data = f.read()
            
            detected = chardet.detect(raw_data)
            encoding = detected.get('encoding', 'utf-8')
            
            # Fallback to utf-8 if encoding is None
            if not encoding:
                encoding = 'utf-8'
                
            text = raw_data.decode(encoding, errors='replace')
            
            if not text.strip():
                return {"success": False, "error": "Could not extract text from TXT. File might be empty."}
                
            return {
                "success": True,
                "text": text,
                "encoding": encoding
            }
        except Exception as e:
            return {"success": False, "error": str(e)}
    
    async def extract_file(self, file_path: str, file_type: str) -> Dict:
        """Route extraction by file type"""
        if file_type.lower() == "pdf":
            return await self.extract_pdf(file_path)
        elif file_type.lower() == "docx":
            return await self.extract_docx(file_path)
        elif file_type.lower() == "txt":
            return await self.extract_txt(file_path)
        else:
            return {"success": False, "error": f"Unsupported file type: {file_type}"}
    
    async def process_document(self, file_path: str, file_type: str, 
                              document_title: str = None) -> Dict:
        """Complete document processing"""
        
        # Extract
        extraction_result = await self.extract_file(file_path, file_type)
        
        if not extraction_result.get("success"):
            return {
                "status": "error",
                "error": extraction_result.get("error", "Unknown extraction error"),
                "document_id": str(uuid.uuid4())
            }
        
        raw_text = extraction_result["text"]
        
        # Clean
        cleaned_text = clean_text(raw_text)
        
        # Chunk
        chunks = chunk_text(cleaned_text, self.chunk_size, self.chunk_overlap)
        
        # Stats
        total_words = len(cleaned_text.split())
        total_tokens = sum(c["token_count"] for c in chunks)
        
        return {
            "status": "success",
            "document_id": str(uuid.uuid4()),
            "file_name": Path(file_path).name,
            "file_type": file_type,
            "extraction": {
                "total_pages": extraction_result.get("page_count", 1),
                "total_words": total_words,
                "total_tokens": total_tokens
            },
            "content": {
                "raw_text": cleaned_text,
                "chunks": chunks,
                "metadata": {
                    "title": document_title or Path(file_path).stem,
                    "processed_at": datetime.utcnow().isoformat()
                }
            },
            "next_action": "ready_for_quiz_generation"
        }

# Initialize singleton
doc_service = DocumentService()
